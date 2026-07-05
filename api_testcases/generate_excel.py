"""生成 API 测试用例 Excel 文件"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 定义所有 API 端点和测试用例
test_cases = [
    # ==================== 认证模块 ====================
    ("TC-AUTH-001", "认证", "用户注册-正向", "POST", "/v1/auth/register",
     '{"account": "test_xxx@test.com", "password": "Test@123456", "nickname": "Tester"}',
     "注册成功，返回用户 ID > 0", "P0", "新用户注册"),

    ("TC-AUTH-002", "认证", "用户注册-重复账号", "POST", "/v1/auth/register",
     '{"account": "user@example.com", "password": "Test@123456", "nickname": "Tmp"}',
     "注册失败，返回错误码", "P1", "重复账号容错"),

    ("TC-AUTH-003", "认证", "用户注册-缺少 account", "POST", "/v1/auth/register",
     '{"password": "Test@123456", "nickname": "T"}',
     "返回 200/400/500", "P2", "参数校验"),

    ("TC-AUTH-004", "认证", "用户注册-缺少 password", "POST", "/v1/auth/register",
     '{"account": "test_xxx@test.com", "nickname": "T"}',
     "返回 200/400", "P2", "参数校验"),

    ("TC-AUTH-005", "认证", "用户注册-空请求体", "POST", "/v1/auth/register",
     '{}',
     "返回 200/400/500", "P2", "边界值"),

    ("TC-AUTH-006", "认证", "用户登录-正向", "POST", "/v1/auth/login",
     '{"account": "user@example.com", "password": "123456"}',
     "登录成功，返回 token（长度 > 50）", "P0", "有效凭证"),

    ("TC-AUTH-007", "认证", "用户登录-错误密码", "POST", "/v1/auth/login",
     '{"account": "user@example.com", "password": "WrongPass"}',
     "登录失败，返回错误信息", "P1", "密码错误"),

    ("TC-AUTH-008", "认证", "用户登录-空密码", "POST", "/v1/auth/login",
     '{"account": "user@example.com", "password": ""}',
     "返回 400/401 UNAUTHORIZED", "P2", "空密码"),

    ("TC-AUTH-009", "认证", "用户登录-缺少 account", "POST", "/v1/auth/login",
     '{"password": "123456"}',
     "返回 400", "P2", "参数校验"),

    ("TC-AUTH-010", "认证", "用户登录-缺少 password", "POST", "/v1/auth/login",
     '{"account": "user@example.com"}',
     "返回 400", "P2", "参数校验"),

    ("TC-AUTH-011", "认证", "获取当前用户信息-有 Token", "GET", "/v1/me",
     "无",
     "返回 200，含 id、account 字段", "P1", "鉴权通过"),

    ("TC-AUTH-012", "认证", "获取当前用户信息-无 Token", "GET", "/v1/me",
     "无",
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-AUTH-013", "认证", "获取当前用户信息-伪造 Token", "GET", "/v1/me",
     "Header: Authorization: Bearer invalid.fake.token",
     "返回 400/401 UNAUTHORIZED", "P2", "Token 无效"),

    # ==================== 地址管理 ====================
    ("TC-ADDR-001", "地址", "添加地址-正向", "POST", "/v1/me/addresses",
     '{"receiver": "张三", "phone": "13800138000", "region": "北京 朝阳区", "detail": "建国路88号", "isDefault": 1}',
     "添加成功，code = 200", "P0", "必填字段完整"),

    ("TC-ADDR-002", "地址", "获取地址列表", "GET", "/v1/me/addresses",
     "无",
     "返回 200，data 为列表", "P1", "列表查询"),

    ("TC-ADDR-003", "地址", "添加地址-缺少 receiver", "POST", "/v1/me/addresses",
     '{"phone": "13800138000", "region": "北京", "detail": "xx路"}',
     "返回 200/400/500", "P2", "参数校验"),

    ("TC-ADDR-004", "地址", "添加地址-无 Token", "POST", "/v1/me/addresses",
     '{"receiver": "X", "phone": "13800138000", "region": "北京", "detail": "xx路"}',
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-ADDR-005", "地址", "更新地址-正向", "PUT", "/v1/me/addresses/{id}",
     '{"receiver": "已更新", "phone": "13600999000", "region": "深圳 南山区", "detail": "更新后的地址", "isDefault": 1}',
     "更新成功", "P2", "地址修改"),

    ("TC-ADDR-006", "地址", "更新不存在的地址", "PUT", "/v1/me/addresses/99999",
     '{"receiver": "X", "phone": "13600136000", "region": "北京", "detail": "xx路"}',
     "返回 200/400/404", "P2", "边界值"),

    ("TC-ADDR-007", "地址", "更新地址-无 Token", "PUT", "/v1/me/addresses/1",
     '{"receiver": "X", "phone": "13800138000", "region": "北京", "detail": "xx路"}',
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-ADDR-008", "地址", "删除地址-正向", "DELETE", "/v1/me/addresses/{id}",
     "无",
     "返回 200/400/404", "P2", "地址删除"),

    ("TC-ADDR-009", "地址", "删除地址-无 Token", "DELETE", "/v1/me/addresses/1",
     "无",
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    # ==================== 商品管理 ====================
    ("TC-PROD-001", "商品", "商品列表-默认分页", "GET", "/v1/products",
     "params: size=5",
     "返回 200，data 为列表，含 id/name/price/stock/categoryId", "P0", "分页查询"),

    ("TC-PROD-002", "商品", "商品搜索-关键字", "GET", "/v1/products",
     "params: keyword=LG, size=3",
     "返回 200，结果含关键字", "P1", "关键字搜索"),

    ("TC-PROD-003", "商品", "商品搜索-分类筛选", "GET", "/v1/products",
     "params: category=5, size=3",
     "返回 200，所有商品 categoryId=5", "P1", "分类筛选"),

    ("TC-PROD-004", "商品", "商品列表-分页边界", "GET", "/v1/products",
     "params: page=1, size=2",
     "返回 200，列表长度 <= 2", "P2", "分页边界"),

    ("TC-PROD-005", "商品", "商品列表-超大页码", "GET", "/v1/products",
     "params: page=99999, size=5",
     "返回 200，data = []", "P2", "空结果"),

    ("TC-PROD-006", "商品", "商品搜索-特殊字符", "GET", "/v1/products",
     'params: keyword=<script>alert(1)</script>',
     "返回 200，不报错", "P2", "XSS 防护"),

    ("TC-PROD-007", "商品", "商品列表-无 Token", "GET", "/v1/products",
     "params: size=2",
     "返回 200，公开接口无需 Token", "P2", "公开访问"),

    ("TC-PROD-008", "商品", "商品详情-有效 ID", "GET", "/v1/products/{id}",
     "路径参数: id=1",
     "返回 200，含 id/name/price/description", "P0", "详情查询"),

    ("TC-PROD-009", "商品", "商品详情-不存在的 ID", "GET", "/v1/products/99999",
     "路径参数: id=99999",
     "返回 200/400/404", "P2", "边界值"),

    ("TC-PROD-010", "商品", "商品详情-非数字 ID", "GET", "/v1/products/abc",
     "路径参数: id=abc",
     "返回 200/400/404/500", "P2", "类型校验"),

    # ==================== 购物车 ====================
    ("TC-CART-001", "购物车", "加入购物车-正向", "POST", "/v1/cart/items",
     '{"productId": 1, "quantity": 1}',
     "返回 200", "P0", "加购"),

    ("TC-CART-002", "购物车", "加入购物车-多数量", "POST", "/v1/cart/items",
     '{"productId": 1, "quantity": 3}',
     "返回 200", "P1", "数量 > 1"),

    ("TC-CART-003", "购物车", "加入购物车-缺少 productId", "POST", "/v1/cart/items",
     '{"quantity": 1}',
     "返回 200/400/500", "P2", "参数校验"),

    ("TC-CART-004", "购物车", "加入购物车-quantity=0", "POST", "/v1/cart/items",
     '{"productId": 1, "quantity": 0}',
     "返回 200/400", "P2", "边界值"),

    ("TC-CART-005", "购物车", "加入购物车-quantity 负数", "POST", "/v1/cart/items",
     '{"productId": 1, "quantity": -1}',
     "返回 200/400", "P2", "边界值"),

    ("TC-CART-006", "购物车", "加入购物车-不存在的商品", "POST", "/v1/cart/items",
     '{"productId": 99999, "quantity": 1}',
     "返回 200/400/404", "P2", "商品不存在"),

    ("TC-CART-007", "购物车", "加入购物车-无 Token", "POST", "/v1/cart/items",
     '{"productId": 1, "quantity": 1}',
     "返回 400/401 UNAUTHORIZED", "P1", "鉴权失败"),

    ("TC-CART-008", "购物车", "获取购物车列表", "GET", "/v1/cart",
     "无",
     "返回 200", "P1", "列表查询"),

    ("TC-CART-009", "购物车", "获取购物车-无 Token", "GET", "/v1/cart",
     "无",
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-CART-010", "购物车", "更新购物车数量", "PUT", "/v1/cart/items/{id}",
     '{"quantity": 5}',
     "返回 200/400/500", "P2", "数量修改"),

    ("TC-CART-011", "购物车", "更新购物车-无 Token", "PUT", "/v1/cart/items/1",
     '{"quantity": 3}',
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-CART-012", "购物车", "删除购物车项", "DELETE", "/v1/cart/items/{id}",
     "无",
     "返回 200/400/404", "P2", "删除"),

    ("TC-CART-013", "购物车", "删除购物车-无 Token", "DELETE", "/v1/cart/items/1",
     "无",
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    # ==================== 订单管理 ====================
    ("TC-ORDER-001", "订单", "下单结算-正向", "POST", "/v1/orders/checkout",
     '{"addressId": 1}',
     "返回 200，含订单 ID", "P0", "正常下单"),

    ("TC-ORDER-002", "订单", "下单-缺少地址", "POST", "/v1/orders/checkout",
     '{}',
     "返回 200/400", "P1", "参数校验"),

    ("TC-ORDER-003", "订单", "下单-无 Token", "POST", "/v1/orders/checkout",
     '{"addressId": 1}',
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-ORDER-004", "订单", "订单列表", "GET", "/v1/orders",
     "params: size=5",
     "返回 200，data 为列表", "P0", "列表查询"),

    ("TC-ORDER-005", "订单", "订单列表-无 Token", "GET", "/v1/orders",
     "params: size=5",
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-ORDER-006", "订单", "订单详情-有效 ID", "GET", "/v1/orders/{id}",
     "路径参数: id=1",
     "返回 200，含订单明细", "P2", "详情查询"),

    ("TC-ORDER-007", "订单", "取消订单-正向", "POST", "/v1/orders/{id}/cancel",
     "路径参数: id=1",
     "返回 200/400", "P1", "订单取消"),

    # ==================== 支付管理 ====================
    ("TC-PAY-001", "支付", "发起支付-正向", "POST", "/v1/payments/{orderId}/pay",
     '{"channel": "alipay"}',
     "返回 200/400/500", "P1", "发起支付"),

    ("TC-PAY-002", "支付", "查询支付状态", "GET", "/v1/payments/{orderId}/status",
     "路径参数: orderId=1",
     "返回 200/400/404", "P1", "状态查询"),

    ("TC-PAY-003", "支付", "支付回调（无需 Token）", "POST", "/v1/payments/webhook",
     '{"orderId": 1, "tradeNo": "T_TEST", "channel": "alipay", "amount": 100, "status": "SUCCESS"}',
     "返回 200/400/500", "P2", "回调通知"),

    ("TC-PAY-004", "支付", "支付不存在的订单", "POST", "/v1/payments/99999/pay",
     '{"channel": "alipay"}',
     "返回 200/400/404", "P2", "边界值"),

    ("TC-PAY-005", "支付", "发起支付-无 Token", "POST", "/v1/payments/1/pay",
     '{"channel": "alipay"}',
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    # ==================== 评价管理 ====================
    ("TC-REV-001", "评价", "创建评价-正向", "POST", "/v1/reviews",
     '{"productId": 1, "content": "非常好用！", "rating": 5}',
     "返回 200/400/500", "P1", "创建评价"),

    ("TC-REV-002", "评价", "创建评价-不存在的商品", "POST", "/v1/reviews",
     '{"productId": 99999, "content": "测试", "rating": 3}',
     "返回 200/400/404", "P2", "边界值"),

    ("TC-REV-003", "评价", "创建评价-缺少必填字段", "POST", "/v1/reviews",
     '{}',
     "返回 200/400", "P2", "参数校验"),

    ("TC-REV-004", "评价", "创建评价-无 Token", "POST", "/v1/reviews",
     '{"productId": 1, "content": "测试", "rating": 5}',
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-REV-005", "评价", "获取评价列表", "GET", "/v1/reviews",
     "params: size=5",
     "返回 200，data 为列表", "P1", "列表查询"),

    ("TC-REV-006", "评价", "按商品 ID 过滤评价", "GET", "/v1/reviews",
     "params: productId=1, size=5",
     "返回 200", "P2", "过滤查询"),

    # ==================== 优惠券 ====================
    ("TC-COUP-001", "优惠券", "获取可用优惠券列表", "GET", "/v1/coupons/available",
     "无",
     "返回 200，含 data 字段", "P1", "列表查询"),

    ("TC-COUP-002", "优惠券", "获取优惠券-无 Token", "GET", "/v1/coupons/available",
     "无",
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-COUP-003", "优惠券", "校验优惠券码-有效", "POST", "/v1/coupons/NEW300/check",
     '{"amount": 6000}',
     "返回 200，含可用状态", "P1", "优惠券校验"),

    ("TC-COUP-004", "优惠券", "校验不存在的优惠券码", "POST", "/v1/coupons/INVALID123/check",
     '{"amount": 1000}',
     "返回 200/400/404", "P2", "边界值"),

    ("TC-COUP-005", "优惠券", "校验优惠券-无 Token", "POST", "/v1/coupons/NEW300/check",
     '{"amount": 1000}',
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    # ==================== 售后管理 ====================
    ("TC-AFTER-001", "售后", "申请售后-正向", "POST", "/v1/aftersales/apply",
     '{"orderId": 1, "reason": "商品质量问题", "type": "REFUND", "amount": 100.00}',
     "返回 200/400/500", "P1", "提交售后"),

    ("TC-AFTER-002", "售后", "申请售后-缺少字段", "POST", "/v1/aftersales/apply",
     '{}',
     "返回 200/400/500", "P2", "参数校验"),

    ("TC-AFTER-003", "售后", "申请售后-无 Token", "POST", "/v1/aftersales/apply",
     '{"orderId": 1, "reason": "测试", "type": "REFUND"}',
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-AFTER-004", "售后", "获取售后列表", "GET", "/v1/aftersales",
     "params: size=5",
     "返回 200", "P1", "列表查询"),

    ("TC-AFTER-005", "售后", "获取售后列表-无 Token", "GET", "/v1/aftersales",
     "params: size=5",
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-AFTER-006", "售后", "取消售后申请", "POST", "/v1/aftersales/{id}/cancel",
     "路径参数: id=1",
     "返回 200/400/404", "P2", "取消售后"),

    # ==================== 通知管理 ====================
    ("TC-NOTIF-001", "通知", "获取通知列表", "GET", "/v1/notifications",
     "params: size=5",
     "返回 200", "P1", "列表查询"),

    ("TC-NOTIF-002", "通知", "获取通知-无 Token", "GET", "/v1/notifications",
     "params: size=5",
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    ("TC-NOTIF-003", "通知", "创建通知", "POST", "/v1/notifications",
     '{"title": "系统通知", "content": "订单已发货", "type": "ORDER"}',
     "返回 200/400/500", "P2", "创建通知"),

    ("TC-NOTIF-004", "通知", "标记单条已读", "POST", "/v1/notifications/{id}/read",
     "路径参数: id=1",
     "返回 200/400/404", "P2", "状态更新"),

    ("TC-NOTIF-005", "通知", "标记全部已读", "POST", "/v1/notifications/markAllRead",
     "无",
     "返回 200/400/500", "P2", "批量更新"),

    ("TC-NOTIF-006", "通知", "清空通知", "DELETE", "/v1/notifications",
     "无",
     "返回 200/400/500", "P2", "批量删除"),

    ("TC-NOTIF-007", "通知", "清空通知-无 Token", "DELETE", "/v1/notifications",
     "无",
     "返回 400/401 UNAUTHORIZED", "P2", "鉴权失败"),

    # ==================== 安全测试 ====================
    ("TC-SEC-001", "安全", "商品搜索-XSS 脚本", "GET", "/v1/products",
     'params: keyword=<script>alert("xss")</script>',
     "返回 200，不执行脚本", "P3", "XSS 防护"),

    ("TC-SEC-002", "安全", "注册账号-XSS 昵称", "POST", "/v1/auth/register",
     '{"account": "xss_xxx@test.com", "password": "Test@123", "nickname": "<img src=x onerror=alert(1)>"}',
     "返回 200/400", "P3", "XSS 防护"),

    ("TC-SEC-003", "安全", "评价内容-XSS 脚本", "POST", "/v1/reviews",
     '{"productId": 1, "content": "<script>alert(1)</script>", "rating": 1}',
     "返回 200/400/500", "P3", "XSS 防护"),

    ("TC-SEC-004", "安全", "商品搜索-SQL 注入", "GET", "/v1/products",
     'params: keyword=\' OR \'1\'=\'1',
     "返回 200", "P3", "SQL 注入防护"),

    ("TC-SEC-005", "安全", "登录接口-SQL 注入", "POST", "/v1/auth/login",
     '{"account": "\' OR \'1\'=\'1", "password": "\' OR \'1\'=\'1"}',
     "返回 200/400", "P3", "SQL 注入防护"),

    ("TC-SEC-006", "安全", "超长关键字搜索", "GET", "/v1/products",
     "params: keyword=A*5000",
     "返回 200/400/414", "P3", "边界值"),

    ("TC-SEC-007", "安全", "超大分页参数", "GET", "/v1/products",
     "params: page=999999, size=1000",
     "返回 200", "P3", "边界值"),

    ("TC-SEC-008", "安全", "超长密码注册", "POST", "/v1/auth/register",
     '{"account": "longpwd_xxx@test.com", "password": "A*200", "nickname": "LongPwdTest"}',
     "返回 200/400", "P3", "边界值"),

    ("TC-SEC-009", "安全", "购物车超大数量", "POST", "/v1/cart/items",
     '{"productId": 1, "quantity": 99999}',
     "返回 200/400", "P3", "边界值"),

    # ==================== 根路径/服务状态 ====================
    ("TC-ROOT-001", "系统", "根路径服务状态", "GET", "/",
     "无",
     "返回 200", "P0", "服务可用性"),

    # ==================== 冒烟测试 Smoke ====================
    ("TC-SMOKE-001", "冒烟", "根路径健康检查", "GET", "/",
     "无",
     "返回 HTTP 200", "P0", "smoke: 服务可用性"),

    ("TC-SMOKE-002", "冒烟", "注册接口可达", "POST", "/v1/auth/register",
     '{"account": "smoke_xxx@test.com", "password": "Smoke@123", "nickname": "SmokeTest"}',
     "返回 200，data.id > 0", "P0", "smoke: 核心接口可达"),

    ("TC-SMOKE-003", "冒烟", "登录接口可达", "POST", "/v1/auth/login",
     '{"account": "smoke_xxx@test.com", "password": "Smoke@123"}',
     "返回 200，token 非空且长度 > 50", "P0", "smoke: 认证可用"),

    ("TC-SMOKE-004", "冒烟", "商品列表接口可达", "GET", "/v1/products",
     "params: size=2",
     "返回 200，data 为列表", "P0", "smoke: 商品服务可达"),

    ("TC-SMOKE-005", "冒烟", "加购接口可达", "POST", "/v1/cart/items",
     '{"productId": 1, "quantity": 1}',
     "返回 200", "P0", "smoke: 购物车服务可达"),

    ("TC-SMOKE-006", "冒烟", "下单接口可达", "POST", "/v1/orders/checkout",
     '{"addressId": 1}',
     "返回 200/400/500", "P0", "smoke: 订单服务可达"),
]


def create_excel(filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "API测试用例"

    # 表头
    headers = [
        "用例ID", "模块", "接口名称", "请求方法", "接口地址",
        "请求参数/Body", "预期结果", "优先级", "备注"
    ]

    # 表头样式
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    # 数据行
    priority_colors = {
        "P0": "FFD7D7",  # 浅红
        "P1": "FFE5B4",  # 浅橙
        "P2": "FFFACD",  # 浅黄
        "P3": "D5F5E3",  # 浅绿
    }

    for row_idx, tc in enumerate(test_cases, 2):
        for col_idx, value in enumerate(tc, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin_border

        # 优先级着色
        priority = tc[7]
        if priority in priority_colors:
            fill = PatternFill(start_color=priority_colors[priority],
                               end_color=priority_colors[priority], fill_type="solid")
            ws.cell(row=row_idx, column=8).fill = fill

    # 列宽
    col_widths = [14, 10, 22, 8, 28, 45, 35, 8, 20]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # 冻结首行
    ws.freeze_panes = "A2"

    # 自动筛选
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(test_cases) + 1}"

    wb.save(filepath)
    print(f"[OK] 已生成: {filepath}")
    print(f"   用例总数: {len(test_cases)}")
    print(f"   P0: {sum(1 for tc in test_cases if tc[7] == 'P0')}")
    print(f"   P1: {sum(1 for tc in test_cases if tc[7] == 'P1')}")
    print(f"   P2: {sum(1 for tc in test_cases if tc[7] == 'P2')}")
    print(f"   P3: {sum(1 for tc in test_cases if tc[7] == 'P3')}")


if __name__ == "__main__":
    output = r"E:\AI\LearningSpace\web-wenrui\api_testcases\ProjectKu_API_TestCases.xlsx"
    create_excel(output)
